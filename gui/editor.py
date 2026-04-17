"""
ProjectEditorApp – Zentrale Applikationsklasse (Dear PyGui)
============================================================

Hält das Pydantic-Project-Modell als Single Source of Truth und verbindet
GUI-Events mit den lib/-Funktionen.  Keine lib/-Logik wird dupliziert.
"""

import json
import sys
import tempfile
from pathlib import Path
from typing import Optional

sys.path.insert(0, str(Path(__file__).parent.parent))

from lib.models import (
    load_project,
    load_project_from_dict,
    save_project,
    Project,
    CPMResult,
)
from lib.models.tasks import SimpleTask, LoopTask, SubTask
from gui.gui_config import load_gui_config

_CFG = load_gui_config()
_COLORS = _CFG.section("colors")


# ---------------------------------------------------------------------------
# Hilfs-Funktionen: Tasks ↔ Tabellenzeilen
# ---------------------------------------------------------------------------

def _parse_id(val: str):
    """Konvertiert String-ID in int, wenn möglich."""
    try:
        return int(val)
    except (ValueError, TypeError):
        return val


def _successors_to_str(task) -> str:
    """Kodiert alle Nachfolger-Listen als 'id[TYP], ...' String."""
    parts = []
    for s in task.successors or []:
        parts.append(f"{s}[EA]")
    for s in task.successors_aa or []:
        parts.append(f"{s}[AA]")
    for s in task.successors_ee or []:
        parts.append(f"{s}[EE]")
    for s in task.successors_ae or []:
        parts.append(f"{s}[AE]")
    return ", ".join(parts)


def _str_to_successors(s: str) -> tuple:
    """Parst 'id[TYP], ...' String in vier Nachfolger-Listen."""
    ea, aa, ee, ae = [], [], [], []
    if not s.strip():
        return ea, aa, ee, ae
    for item in s.split(","):
        item = item.strip()
        if not item:
            continue
        if "[" in item:
            id_part, typ = item.rsplit("[", 1)
            typ = typ.rstrip("]").upper()
        else:
            id_part, typ = item, "EA"
        id_val = _parse_id(id_part.strip())
        if typ == "AA":
            aa.append(id_val)
        elif typ == "EE":
            ee.append(id_val)
        elif typ == "AE":
            ae.append(id_val)
        else:
            ea.append(id_val)
    return ea, aa, ee, ae


def tasks_to_rows(tasks: list) -> list:
    """Flacht Project.tasks (inkl. Subtasks) in Tabellenzeilen ab."""
    rows = []
    for task in tasks:
        if isinstance(task, LoopTask):
            rows.append({
                "row_type": "loop",
                "id": str(task.id),
                "name": task.name,
                "duration": "",
                "resources": "",
                "successors_str": _successors_to_str(task),
                "cost": str(task.cost) if task.cost is not None else "",
                "loop_until": task.loop_until or "",
                "cycle_prefix": task.cycle_prefix or "F",
                "loop_count": str(task.loop_count) if task.loop_count is not None else "",
                "volume_per_cycle": str(task.volume_per_cycle) if task.volume_per_cycle is not None else "",
                "parent_id": "",
            })
            for sub in (task.subtasks or []):
                rows.append({
                    "row_type": "subtask",
                    "id": sub.id or "",
                    "name": sub.name,
                    "duration": sub.duration or sub.duration_formula or "",
                    "resources": ", ".join(sub.resources or []),
                    "successors_str": "",
                    "cost": "",
                    "loop_until": "",
                    "cycle_prefix": "",
                    "loop_count": "",
                    "volume_per_cycle": "",
                    "parent_id": str(task.id),
                })
        else:
            rows.append({
                "row_type": "simple",
                "id": str(task.id),
                "name": task.name,
                "duration": task.duration or "",
                "resources": ", ".join(task.resources or []),
                "successors_str": _successors_to_str(task),
                "cost": str(task.cost) if task.cost is not None else "",
                "loop_until": "",
                "cycle_prefix": "",
                "loop_count": "",
                "volume_per_cycle": "",
                "parent_id": "",
            })
    return rows


def rows_to_tasks(rows: list) -> list:
    """Rekonstruiert Project.tasks aus flachen Tabellenzeilen."""
    tasks = []
    i = 0
    while i < len(rows):
        row = rows[i]
        if row.get("row_type") == "loop":
            subtasks = []
            j = i + 1
            while (
                j < len(rows)
                and rows[j].get("row_type") == "subtask"
                and rows[j].get("parent_id") == row["id"]
            ):
                sr = rows[j]
                res = [r.strip() for r in sr.get("resources", "").split(",") if r.strip()]
                sub = SubTask(
                    id=sr["id"] or None,
                    name=sr["name"],
                    resources=res or None,
                    duration=sr.get("duration") or None,
                )
                subtasks.append(sub)
                j += 1
            ea, aa, ee, ae = _str_to_successors(row.get("successors_str", ""))
            lc = int(row["loop_count"]) if row.get("loop_count") else None
            vpc = float(row["volume_per_cycle"]) if row.get("volume_per_cycle") else None
            task = LoopTask(
                id=_parse_id(row["id"]),
                name=row["name"],
                loop_until=row.get("loop_until") or "total_volume <= 0",
                successors=ea,
                successors_aa=aa,
                successors_ee=ee,
                successors_ae=ae,
                cycle_prefix=row.get("cycle_prefix") or "F",
                subtasks=subtasks,
                loop_count=lc,
                volume_per_cycle=vpc,
            )
            tasks.append(task)
            i = j
        elif row.get("row_type") == "subtask":
            i += 1  # Verwaiste Subtasks überspringen
        else:
            ea, aa, ee, ae = _str_to_successors(row.get("successors_str", ""))
            cost_str = row.get("cost", "")
            cost = float(cost_str) if cost_str else None
            res = [r.strip() for r in row.get("resources", "").split(",") if r.strip()]
            task = SimpleTask(
                id=_parse_id(row["id"]),
                name=row["name"],
                duration=row.get("duration") or "0d",
                successors=ea,
                successors_aa=aa,
                successors_ee=ee,
                successors_ae=ae,
                resources=res or None,
                cost=cost,
            )
            tasks.append(task)
            i += 1
    return tasks


# ---------------------------------------------------------------------------
# Hauptklasse
# ---------------------------------------------------------------------------

class ProjectEditorApp:
    """
    Kapselt den gesamten GUI-State und alle Event-Handler.
    Importiert lib/ direkt; keine Logik wird dupliziert.
    """

    def __init__(self):
        self.project: Optional[Project] = None
        self.active_sections: dict = {
            "resources": False,
            "persons": False,
            "resting_times": False,
        }
        self.last_cpm_result: Optional[CPMResult] = None
        self.dirty: bool = False
        self._task_rows: list = []
        self._last_save_path: Optional[Path] = None

    # ------------------------------------------------------------------
    # Datei-Operationen
    # ------------------------------------------------------------------

    def load_from_file(self, file_path: Path) -> None:
        try:
            self.project = load_project(file_path)
            self._task_rows = tasks_to_rows(self.project.tasks)
            self._detect_active_sections()
            self._last_save_path = file_path
            self.last_cpm_result = None
            self.refresh_all()
            self._set_status(f"Projekt geladen: {self.project.project}")
        except Exception as e:
            self._set_status(f"Fehler beim Laden: {e}")

    def load_from_bytes(self, content: bytes, filename: str = "") -> None:
        try:
            data = json.loads(content.decode("utf-8"))
            self.project = load_project_from_dict(data)
            self._task_rows = tasks_to_rows(self.project.tasks)
            self._detect_active_sections()
            self.refresh_all()
            self._set_status(f"Projekt geladen: {self.project.project}")
        except Exception as e:
            self._set_status(f"Fehler beim Laden: {e}")

    def save_to_file(self, file_path: Optional[Path] = None) -> Path:
        """Serialisiert den aktuellen State als sproject-JSON."""
        self._sync_to_model()
        if file_path is None:
            name = (self.project.project or "project").replace(" ", "_")
            file_path = Path.cwd() / f"{name}.json"
        save_project(self.project, file_path)
        self._last_save_path = file_path
        self.dirty = False
        self._set_status(f"Gespeichert: {file_path}")
        return file_path

    def save_quick(self) -> None:
        """Speichert ohne Dialog: gleicher Pfad wie beim letzten Laden/Speichern."""
        if self.project is None:
            self._set_status("Kein Projekt geladen.")
            return
        self.save_to_file(self._last_save_path)

    def get_json_bytes(self) -> bytes:
        """Gibt das aktuelle Projekt als JSON-Bytes zurück."""
        self._sync_to_model()
        tmp = Path(tempfile.mktemp(suffix=".json"))
        save_project(self.project, tmp)
        data = tmp.read_bytes()
        try:
            tmp.unlink()
        except OSError:
            pass
        return data

    # ------------------------------------------------------------------
    # Berechnung
    # ------------------------------------------------------------------

    def run_calculation(self) -> None:
        if self.project is None:
            self._set_status("Kein Projekt geladen.")
            return
        try:
            self._sync_to_model()
            result = self.project.calculate_cpm()
            self.last_cpm_result = result
            self._show_result(result)
            self._set_status(
                f"Berechnung abgeschlossen – Dauer: {result.project_duration}"
            )
        except Exception as e:
            self._set_status(f"Fehler bei der Berechnung: {e}")

    # ------------------------------------------------------------------
    # Interne Hilfsmethoden
    # ------------------------------------------------------------------

    def _detect_active_sections(self) -> None:
        if self.project is None:
            return
        self.active_sections["resources"] = bool(self.project.resources)
        self.active_sections["persons"] = bool(self.project.persons)
        self.active_sections["resting_times"] = bool(self.project.resting_times)

    def _read_rows_from_ui(self) -> None:
        """Liest aktuelle Eingabewerte der Task-Tabelle in _task_rows zurück."""
        import dearpygui.dearpygui as dpg
        for i, row in enumerate(self._task_rows):
            for field in ("id", "name", "duration", "successors_str", "resources", "cost"):
                tag = f"task_{i}_{field}"
                if dpg.does_item_exist(tag):
                    row[field] = dpg.get_value(tag) or ""

    def _sync_to_model(self) -> None:
        """Überträgt GUI-State ins Pydantic-Modell."""
        import dearpygui.dearpygui as dpg
        if self.project is None:
            return

        self._read_rows_from_ui()
        self.project.tasks = rows_to_tasks(self._task_rows)

        def _get(tag: str) -> str:
            try:
                return dpg.get_value(tag) or ""
            except Exception:
                return ""

        v = _get("inp_project_name")
        if v:
            self.project.project = v
        v = _get("inp_project_start")
        if v:
            self.project.project_start = v
        v = _get("inp_unit")
        if v:
            self.project.unit = v
        v = _get("inp_total_volume")
        if v:
            try:
                self.project.total_volume = float(v)
            except ValueError:
                pass
        v = _get("inp_total_hours")
        if v:
            try:
                self.project.total_hours = int(v)
            except ValueError:
                pass
        v = _get("inp_order_volume")
        if v:
            try:
                self.project.order_volume = int(v)
            except ValueError:
                pass

    def refresh_all(self) -> None:
        """Aktualisiert alle UI-Komponenten nach dem Laden eines Projekts."""
        import dearpygui.dearpygui as dpg
        if self.project is None:
            return

        def _set(tag: str, val: str) -> None:
            try:
                if dpg.does_item_exist(tag):
                    dpg.set_value(tag, val)
            except Exception:
                pass

        _set("inp_project_name", self.project.project or "")
        _set("inp_project_start", self.project.project_start or "")
        _set("inp_unit", self.project.unit or "days")
        _set(
            "inp_total_hours",
            str(self.project.total_hours) if self.project.total_hours is not None else "",
        )
        _set(
            "inp_total_volume",
            str(self.project.total_volume) if self.project.total_volume is not None else "",
        )
        _set(
            "inp_order_volume",
            str(self.project.order_volume) if self.project.order_volume is not None else "",
        )

        from gui.components.task_table import rebuild_task_table
        rebuild_task_table(self)

        self._refresh_resources()
        self._refresh_persons()
        self._refresh_resting_times()
        self._update_section_visibility()
        self._update_sidebar()

        # Ergebnis-Panel beim Laden immer ausblenden und leeren
        try:
            if dpg.does_item_exist("result_container"):
                dpg.configure_item("result_container", show=False)
                dpg.delete_item("result_container", children_only=True)
        except Exception:
            pass

    def _update_section_visibility(self) -> None:
        import dearpygui.dearpygui as dpg
        for section, tag in [
            ("resources", "section_resources"),
            ("persons", "section_persons"),
            ("resting_times", "section_resting_times"),
        ]:
            is_active = self.active_sections.get(section, False)
            try:
                if dpg.does_item_exist(tag):
                    dpg.configure_item(tag, show=is_active)
            except Exception:
                pass

    def _update_sidebar(self) -> None:
        import dearpygui.dearpygui as dpg
        for section, toggle_tag in [
            ("resources", "sidebar_res_toggle"),
            ("persons", "sidebar_pers_toggle"),
            ("resting_times", "sidebar_rest_toggle"),
        ]:
            is_active = self.active_sections.get(section, False)
            try:
                if dpg.does_item_exist(toggle_tag):
                    dpg.set_value(toggle_tag, is_active)
            except Exception:
                pass

    def toggle_section(self, section: str) -> None:
        self.active_sections[section] = not self.active_sections.get(section, False)
        self._update_section_visibility()
        self._update_sidebar()

    def add_task_row(self, is_loop: bool = False) -> None:
        """Fügt eine neue Zeile zur Task-Tabelle hinzu."""
        self._read_rows_from_ui()
        numeric_ids = [
            int(r["id"]) for r in self._task_rows if str(r.get("id", "")).isdigit()
        ]
        new_id = str(max(numeric_ids or [0]) + 1)
        if is_loop:
            row = {
                "row_type": "loop",
                "id": new_id,
                "name": "Neuer Loop-Task",
                "duration": "",
                "resources": "",
                "successors_str": "",
                "cost": "",
                "loop_until": "total_volume <= 0",
                "cycle_prefix": "F",
                "loop_count": "",
                "volume_per_cycle": "",
                "parent_id": "",
            }
        else:
            row = {
                "row_type": "simple",
                "id": new_id,
                "name": "Neuer Task",
                "duration": "1d",
                "resources": "",
                "successors_str": "",
                "cost": "",
                "loop_until": "",
                "cycle_prefix": "",
                "loop_count": "",
                "volume_per_cycle": "",
                "parent_id": "",
            }
        self._task_rows.append(row)
        self.dirty = True
        from gui.components.task_table import rebuild_task_table
        rebuild_task_table(self)

    def delete_task_row(self, row_idx: int) -> None:
        """Löscht eine Zeile (und zugehörige Subtasks) aus der Task-Tabelle."""
        self._read_rows_from_ui()
        if row_idx >= len(self._task_rows):
            return
        row = self._task_rows[row_idx]
        indices_to_remove = [row_idx]
        if row.get("row_type") == "loop":
            loop_id = row["id"]
            for j, r in enumerate(self._task_rows):
                if r.get("parent_id") == loop_id:
                    indices_to_remove.append(j)
        for idx in sorted(set(indices_to_remove), reverse=True):
            self._task_rows.pop(idx)
        self.dirty = True
        from gui.components.task_table import rebuild_task_table
        rebuild_task_table(self)

    def add_subtask_to_loop(self, loop_row_idx: int) -> None:
        """Fügt einem Loop-Task einen neuen Subtask hinzu."""
        self._read_rows_from_ui()
        if loop_row_idx >= len(self._task_rows):
            return
        loop_id = self._task_rows[loop_row_idx]["id"]
        insert_pos = loop_row_idx + 1
        for j in range(loop_row_idx + 1, len(self._task_rows)):
            if self._task_rows[j].get("parent_id") == loop_id:
                insert_pos = j + 1
            else:
                break
        self._task_rows.insert(insert_pos, {
            "row_type": "subtask",
            "id": "",
            "name": "Neuer Subtask",
            "duration": "",
            "resources": "",
            "successors_str": "",
            "cost": "",
            "loop_until": "",
            "cycle_prefix": "",
            "loop_count": "",
            "volume_per_cycle": "",
            "parent_id": loop_id,
        })
        self.dirty = True
        from gui.components.task_table import rebuild_task_table
        rebuild_task_table(self)

    # ------------------------------------------------------------------
    # Abschnitts-Inhalte neu aufbauen
    # ------------------------------------------------------------------

    def _refresh_resources(self) -> None:
        import dearpygui.dearpygui as dpg
        if not dpg.does_item_exist("resources_content"):
            return
        dpg.delete_item("resources_content", children_only=True)
        if not self.project or not self.project.resources:
            dpg.add_text("Keine Ressourcen definiert.", parent="resources_content")
            return
        with dpg.table(
            parent="resources_content",
            header_row=True,
            resizable=True,
            borders_outerH=True,
            borders_innerV=True,
            borders_outerV=True,
            row_background=True,
        ):
            _rt = _CFG.section("resources_table")
            dpg.add_table_column(label="ID", width_fixed=True, init_width_or_weight=_rt.get("col_id_width", 100))
            dpg.add_table_column(label="Name", width_stretch=True)
            dpg.add_table_column(label="Typ", width_fixed=True, init_width_or_weight=_rt.get("col_type_width", 90))
            dpg.add_table_column(label="Farbe", width_fixed=True, init_width_or_weight=_rt.get("col_color_width", 80))
            for r in self.project.resources:
                with dpg.table_row():
                    dpg.add_text(r.id)
                    dpg.add_text(r.name or "")
                    dpg.add_text(r.type or "")
                    color = r.color or ""
                    dpg.add_text(f"#{color}" if color else "–")

    def _refresh_persons(self) -> None:
        import dearpygui.dearpygui as dpg
        if not dpg.does_item_exist("persons_content"):
            return
        dpg.delete_item("persons_content", children_only=True)
        if not self.project or not self.project.persons:
            dpg.add_text("Keine Personen definiert.", parent="persons_content")
            return
        with dpg.table(
            parent="persons_content",
            header_row=True,
            resizable=True,
            borders_outerH=True,
            borders_innerV=True,
            borders_outerV=True,
            row_background=True,
        ):
            _pt = _CFG.section("persons_table")
            dpg.add_table_column(label="ID", width_fixed=True, init_width_or_weight=_pt.get("col_id_width", 90))
            dpg.add_table_column(label="Name", width_stretch=True)
            dpg.add_table_column(label="Rolle", width_stretch=True)
            dpg.add_table_column(label="EUR/h", width_fixed=True, init_width_or_weight=_pt.get("col_rate_width", 70))
            dpg.add_table_column(label="Info", width_stretch=True)
            for p in self.project.persons:
                with dpg.table_row():
                    dpg.add_text(p.id)
                    dpg.add_text(p.name)
                    dpg.add_text(p.role)
                    dpg.add_text(str(p.hourly_rate))
                    info = []
                    if p.vacation:
                        info.append(f"Urlaub: {len(p.vacation)}")
                    if p.workinghours_override:
                        info.append("Teilzeit")
                    dpg.add_text(", ".join(info) if info else "–")

    def _refresh_resting_times(self) -> None:
        import dearpygui.dearpygui as dpg
        if not dpg.does_item_exist("resting_times_content"):
            return
        dpg.delete_item("resting_times_content", children_only=True)
        if not self.project or not self.project.resting_times:
            dpg.add_text(
                "Keine Ruhezeitregeln definiert.", parent="resting_times_content"
            )
            return
        with dpg.table(
            parent="resting_times_content",
            header_row=True,
            resizable=True,
            borders_outerH=True,
            borders_innerV=True,
            borders_outerV=True,
        ):
            _rtt = _CFG.section("resting_times_table")
            dpg.add_table_column(label="Nach (h)", width_fixed=True, init_width_or_weight=_rtt.get("col_hours_width", 90))
            dpg.add_table_column(label="Pause", width_fixed=True, init_width_or_weight=_rtt.get("col_pause_width", 80))
            dpg.add_table_column(label="Hinweis", width_stretch=True)
            for ri in self.project.resting_times:
                with dpg.table_row():
                    dpg.add_text(str(ri.after_hours))
                    dpg.add_text(ri.duration)
                    dpg.add_text(ri.note or "–")

    # ------------------------------------------------------------------
    # Ergebnis-Panel
    # ------------------------------------------------------------------

    def _show_result(self, result: CPMResult) -> None:
        import dearpygui.dearpygui as dpg
        if not dpg.does_item_exist("result_container"):
            return
        dpg.configure_item("result_container", show=True)
        dpg.delete_item("result_container", children_only=True)

        dpg.add_text(
            f"Projektdauer: {result.project_duration}",
            parent="result_container",
            color=_COLORS.get("result_duration", (100, 180, 255)),
        )
        crit_str = " -> ".join(str(t) for t in result.critical_path)
        dpg.add_text(
            f"Kritischer Pfad: {crit_str}",
            parent="result_container",
            color=_COLORS.get("result_critical_path", (220, 80, 80)),
        )
        dpg.add_spacer(parent="result_container", height=4)

        # Ergebnis-Tabelle
        crit_indices = []
        visible_idx = 0
        with dpg.table(
            parent="result_container",
            tag="result_table",
            header_row=True,
            resizable=True,
            borders_outerH=True,
            borders_innerV=True,
            borders_outerV=True,
            row_background=True,
        ):
            _rest = _CFG.section("result_table")
            dpg.add_table_column(label="#", width_fixed=True, init_width_or_weight=_rest.get("col_id_width", 60))
            dpg.add_table_column(label="Name", width_stretch=True)
            dpg.add_table_column(label="FAZ", width_fixed=True, init_width_or_weight=_rest.get("col_faz_width", 70))
            dpg.add_table_column(label="FEZ", width_fixed=True, init_width_or_weight=_rest.get("col_fez_width", 70))
            dpg.add_table_column(label="SAZ", width_fixed=True, init_width_or_weight=_rest.get("col_saz_width", 70))
            dpg.add_table_column(label="SEZ", width_fixed=True, init_width_or_weight=_rest.get("col_sez_width", 70))
            dpg.add_table_column(label="Puffer", width_fixed=True, init_width_or_weight=_rest.get("col_puffer_width", 70))
            dpg.add_table_column(label="Krit.", width_fixed=True, init_width_or_weight=_rest.get("col_krit_width", 45))

            for tid, t in result.tasks.items():
                if t.is_break:
                    continue
                with dpg.table_row():
                    dpg.add_text(str(tid))
                    dpg.add_text(t.name)
                    dpg.add_text(f"{t.faz:.1f}")
                    dpg.add_text(f"{t.fez:.1f}")
                    dpg.add_text(f"{t.saz:.1f}")
                    dpg.add_text(f"{t.sez:.1f}")
                    dpg.add_text(f"{t.puffer:.1f}")
                    dpg.add_text("★" if t.is_critical else "")
                if t.is_critical:
                    crit_indices.append(visible_idx)
                visible_idx += 1

        for idx in crit_indices:
            dpg.highlight_table_row("result_table", idx, list(_COLORS.get("critical_row", (220, 80, 80, 80))))

        # Export-Buttons
        dpg.add_spacer(parent="result_container", height=4)
        with dpg.group(horizontal=True, parent="result_container"):
            dpg.add_button(label="↓ Markdown", callback=self._export_markdown)
            dpg.add_button(label="↓ Excel", callback=self._export_excel)
            dpg.add_button(label="↓ JSON", callback=self._export_json)

    # ------------------------------------------------------------------
    # Export-Aktionen
    # ------------------------------------------------------------------

    def export_json(self, target_dir: Optional[Path] = None) -> Path:
        """Exportiert das Projekt als JSON in target_dir."""
        if self.project is None:
            raise RuntimeError("Kein Projekt vorhanden.")
        d = target_dir or Path.cwd()
        name = (self.project.project or "project").replace(" ", "_")
        out = d / f"{name}.json"
        self.save_to_file(out)
        return out

    def export_markdown(self, target_dir: Optional[Path] = None) -> Path:
        """Exportiert CPM-Ergebnis als Markdown in target_dir."""
        if not self.last_cpm_result:
            raise RuntimeError("Kein Berechnungsergebnis. Bitte erst berechnen.")
        from lib.markdown_export import export_cpm_to_markdown
        d = target_dir or Path.cwd()
        project_name = self.project.project if self.project else "Projekt"
        out = d / f"{project_name.replace(' ', '_')}.md"
        export_cpm_to_markdown(self.last_cpm_result, out, project_name, self.project)
        return out

    def export_excel(self, target_dir: Optional[Path] = None) -> Path:
        """Exportiert CPM-Ergebnis als Excel in target_dir."""
        if not self.last_cpm_result or not self.project:
            raise RuntimeError("Kein Berechnungsergebnis. Bitte erst berechnen.")
        from openpyxl import Workbook
        from lib.excel_reports import create_gantt_chart, load_excel_export_config
        from lib.models.reports import GanttReport
        d = target_dir or Path.cwd()
        project_name = self.project.project or "project"
        out = d / f"{project_name.replace(' ', '_')}.xlsx"

        cfg_dir = Path(__file__).resolve().parent.parent / "cfg"
        full_config = load_excel_export_config(cfg_dir) if cfg_dir.exists() else {}
        report = GanttReport(
            id='gantt_chart',
            name='Gantt Chart',
            headline='Gantt Chart',
            type='gantt',
            columns=['ID', 'Name', 'Start', 'Ende', 'Aufwand', 'Diagramm'],
            loadunit=self.last_cpm_result.time_unit,
        )

        wb = Workbook()
        create_gantt_chart(wb, self.project, self.last_cpm_result, report, full_config)
        wb.save(out)
        return out

    def export_txt(self, target_dir: Optional[Path] = None) -> Path:
        """Exportiert CPM-Ergebnis als einfache Textdatei in target_dir."""
        if not self.last_cpm_result or not self.project:
            raise RuntimeError("Kein Berechnungsergebnis. Bitte erst berechnen.")
        d = target_dir or Path.cwd()
        result = self.last_cpm_result
        project_name = self.project.project or "Projekt"
        out = d / f"{project_name.replace(' ', '_')}.txt"

        lines = [
            f"Projekt:          {project_name}",
            f"Projektdauer:     {result.project_duration}",
            f"Kritischer Pfad:  {' → '.join(str(t) for t in result.critical_path)}",
            "",
            f"{'#':<6} {'Name':<30} {'FAZ':>7} {'FEZ':>7} {'SAZ':>7} {'SEZ':>7} {'Puffer':>7}  Krit.",
            "-" * 75,
        ]
        for tid, t in result.tasks.items():
            if t.is_break:
                continue
            krit = "★" if t.is_critical else ""
            lines.append(
                f"{str(tid):<6} {t.name:<30} {t.faz:>7.1f} {t.fez:>7.1f}"
                f" {t.saz:>7.1f} {t.sez:>7.1f} {t.puffer:>7.1f}  {krit}"
            )
        out.write_text("\n".join(lines), encoding="utf-8")
        return out

    # Rückwärtskompatible Shortcuts (vom Ergebnis-Panel aufgerufen)
    def _export_json(self)     -> None: self._wrap_export(self.export_json)
    def _export_markdown(self) -> None: self._wrap_export(self.export_markdown)
    def _export_excel(self)    -> None: self._wrap_export(self.export_excel)

    def _get_export_dir(self) -> Path:
        """Gibt das Standard-Exportverzeichnis (PV_RESULTS) zurück und erstellt es bei Bedarf."""
        import os
        d = Path(os.environ.get("PV_RESULTS", "")) if os.environ.get("PV_RESULTS") else Path(__file__).resolve().parent.parent / "results"
        d.mkdir(exist_ok=True)
        return d

    def _wrap_export(self, fn) -> None:
        try:
            out = fn(self._get_export_dir())
            self._set_status(f"Gespeichert: {out}")
        except Exception as e:
            self._set_status(f"Export fehlgeschlagen: {e}")

    # ------------------------------------------------------------------
    # Status
    # ------------------------------------------------------------------

    def _set_status(self, msg: str) -> None:
        import dearpygui.dearpygui as dpg
        try:
            if dpg.does_item_exist("statusbar_text"):
                dpg.set_value("statusbar_text", msg)
        except Exception:
            pass
